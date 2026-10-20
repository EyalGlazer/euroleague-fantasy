import openpyxl, json, sys
# Standings (rank) and next round fixtures, transcribed from the screenshots (home team listed first)
TEAMS = {"PAO":("Panathinaikos",1),"FBT":("Fenerbahçe",2),"OLY":("Olympiacos",3),"VBC":("Valencia Basket",4),"CZV":("Crvena zvezda",5),
 "HTA":("Hapoel Tel Aviv",6),"VIR":("Virtus Bologna",7),"DUB":("Dubai",8),"RMB":("Real Madrid",9),"BAR":("Barcelona",10),
 "BJK":("Beşiktaş",11),"PAR":("Partizan",12),"EFS":("Anadolu Efes",13),"ZAL":("Žalgiris",14),"BAY":("Bayern Munich",15),
 "MTA":("Maccabi Tel Aviv",16),"ASV":("LDLC ASVEL",17),"PBB":("Paris",18),"MIL":("Milano",19),"KBA":("Baskonia",20)}
# Next round (7) fixtures, home first; 1 game in Turn 1 (21/10), 6 in Turn 2 (22/10), 3 in Turn 3 (23/10)
FIXTURES = [("HTA","VIR"),
            ("EFS","BJK"),("BAY","DUB"),("MTA","RMB"),("PAO","KBA"),("MIL","BAR"),("PAR","CZV"),
            ("ZAL","ASV"),("OLY","FBT"),("PBB","VBC")]
# Rounds 1-6 fixtures (home first), transcribed from screenshots. Round 4 last game (BAR-ZAL) was cut off in the
# screenshot and inferred as the two teams not otherwise listed; its home/away order is unknown.
ROUNDS = {  # round number -> fixtures; add a new entry each round
 1: [("DUB","RMB"),("HTA","BAY"),("CZV","ZAL"),("PAO","PBB"),("BAR","EFS"),("KBA","OLY"),("ASV","MTA"),("BJK","VBC"),("FBT","VIR"),("PAR","MIL")],
 2: [("DUB","BAR"),("EFS","RMB"),("ZAL","OLY"),("FBT","BAY"),("CZV","HTA"),("VBC","KBA"),("MIL","VIR"),("PBB","PAR"),("MTA","BJK"),("PAO","ASV")],
 3: [("HTA","RMB"),("CZV","EFS"),("VIR","OLY"),("PBB","ZAL"),("BJK","BAR"),("FBT","DUB"),("ASV","VBC"),("BAY","PAR"),("PAO","MTA"),("KBA","MIL")],
 4: [("PBB","ASV"),("MTA","MIL"),("DUB","CZV"),("BAY","VIR"),("PAO","FBT"),("VBC","HTA"),("RMB","PAR"),("OLY","EFS"),("KBA","BJK"),("BAR","ZAL",None)],
 5: [("FBT","ZAL"),("ASV","CZV"),("KBA","DUB"),("MIL","RMB"),("BAR","MTA"),("PAR","PAO"),("VBC","OLY"),("HTA","BJK"),("EFS","BAY"),("PBB","VIR")],
 6: [("FBT","PAR"),("MIL","DUB"),("VBC","MTA"),("RMB","ASV"),
            ("BJK","BAY"),("ZAL","EFS"),("CZV","PBB"),("PAO","HTA"),("BAR","OLY"),("VIR","KBA")],
}
OPP = {}
for rn_, fx in ROUNDS.items():
    m = {}
    for h, a, *unk in fx:
        known = not unk  # a trailing None marks an unknown home/away order
        m[h] = (a, "H" if known else ""); m[a] = (h, "A" if known else "")
    assert len(m) == 20, "every team must appear exactly once per round"
    OPP[rn_] = m
FIXTURE_TURNS = [1]*1 + [2]*6 + [3]*3  # per FIXTURES entry: 1 game in Turn 1, 6 in Turn 2, 3 in Turn 3
NEXT = {}; TURN = {}
for (h, a), t in zip(FIXTURES, FIXTURE_TURNS):
    NEXT[h] = (a, "H"); NEXT[a] = (h, "A"); TURN[h] = TURN[a] = t
src = sys.argv[1]
import glob
def key(r): return r[0] if isinstance(r[0], int) else f"{r[1]} {r[2]}|{r[5]}"
rounds = {}
for path in sys.argv[2:]:
    w = openpyxl.load_workbook(path, data_only=True).active
    rn_ = next(int(str(r[0]).split(":")[1]) for r in w.iter_rows(values_only=True) if str(r[0]).startswith("Reference matchdays"))
    rounds[rn_] = ({key(r): (float(r[6]), float(r[7])) for r in w.iter_rows(min_row=2, values_only=True) if r[4] is not None})
# hover charts show only the last 5 rounds
LAST = sorted(rounds)[-5:]
ws = openpyxl.load_workbook(src, data_only=True).active
ref = ""
players = []
for r in ws.iter_rows(min_row=2, values_only=True):
    if r[4] is None:
        if r[0]: ref = str(r[0])
        continue
    fpt, price, plus = float(r[6]), float(r[7]), float(r[8])
    players.append(dict(name=f"{r[1]} {r[2]}".strip(), pos=r[4], team=r[5], fpt=fpt, price=price, plus=plus,
                        fptpm=round(fpt/price, 3) if price else 0, pluspm=round(plus/price, 3) if price else 0,
                        rn=LAST, opps=[OPP[n].get(r[5]) for n in LAST], nextopp=TEAMS[NEXT[r[5]][0]][0], nextrank=TEAMS[NEXT[r[5]][0]][1], nextha=NEXT[r[5]][1], nextturn="T%d" % TURN[r[5]],
                        hist=[list(rounds[n][key(r)]) if key(r) in rounds[n] else None for n in LAST]))
import base64
logo = "data:image/png;base64," + base64.b64encode(open("euroleague-logo.png","rb").read()).decode()
html = open("template.html").read().replace("__LOGO__", logo).replace("__DATA__", json.dumps(players)).replace("__REF__", ref)
open("fantasy.html", "w").write(html)
print(len(players), "players")
